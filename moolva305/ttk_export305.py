"""Technology-card PDF and Word templates; fields following GOST 31987-2012.

A text name and system timestamp never become an electronic signature.
Drafts have a prominent status banner and may not be represented as approved.
"""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape


def _common(venue,item):
    data=item["data"]
    is_draft=item["status"]!="approved"
    return data,is_draft,(
        "ПРОЕКТ — технологические параметры не утверждены. НЕ ИСПОЛЬЗОВАТЬ В ПРОИЗВОДСТВЕ."
        if is_draft else
        "Утверждено в учётной системе; необходима надлежащая подпись уполномоченного лица."
    )


def _nut_text(data):
    nut=data.get("nutrition",{})
    def v(x):
        return "не определено" if x in (None,"") else str(x).replace(".",",")
    return "Белки: "+v(nut.get("protein_g"))+" г; жиры: "+v(nut.get("fat_g"))+" г; углеводы: "+v(nut.get("carb_g"))+" г; энергия: "+v(nut.get("kcal"))+" ккал ("+data.get("nutrition_per","на порцию")+")."


def _sections(data):
    return [
      ("1. Область применения",data.get("scope","")),
      ("2. Требования к сырью",data.get("raw_requirements","")),
      ("3. Рецептура",""),
      ("4. Технологический процесс",data.get("technology","")),
      ("5. Оформление, подача, реализация и хранение",
        data.get("presentation","")+"\n"+data.get("storage","")),
      ("6. Показатели качества и безопасности",
        data.get("quality","")+"\n"+data.get("safety","")),
      ("7. Пищевая и энергетическая ценность",
        _nut_text(data)+"\nИсточник и метод: "+data.get("nutrition_basis","")),
      ("8. Аллергены и примечания",
        ("Аллергены: "+(", ".join(data.get("allergens",[])) or "Не подтверждено.")+
          "\n"+data.get("allergen_note","")+"\n"+data.get("author_note","")).strip()),
    ]


def as_pdf(output,venue,item):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,KeepTogether
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.units import mm
    from exports import _register_pdf_font

    data,draft,status=_common(venue,item)
    font=_register_pdf_font()
    style=ParagraphStyle("t305",fontName=font,fontSize=9,leading=13,spaceAfter=6)
    small=ParagraphStyle("sm305",parent=style,fontSize=8,leading=10)
    title=ParagraphStyle("title305",parent=style,fontSize=14,leading=19,
                         alignment=TA_CENTER,spaceAfter=12)
    heading=ParagraphStyle("head305",parent=style,fontSize=10,leading=14,
                           spaceBefore=10,spaceAfter=6)
    def p(s,sty=style):
        return Paragraph(escape(str(s or "—")).replace("\n","<br/>"),sty)
    output=Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    doc=SimpleDocTemplate(str(output),pagesize=A4,
                          rightMargin=16*mm,leftMargin=16*mm,
                          topMargin=15*mm,bottomMargin=15*mm,
                          title=data["title"],author="MoOLVa ХАССП Easy")
    story=[
      p(venue.get("legal_name") or venue.get("venue_name"),small),
      p("Объект: "+str(venue.get("venue_name") or ""),small),
      p("Адрес: "+str(venue.get("address") or "не заполнен"),small),
      p(status,heading),
      p(("ТЕХНИКО-ТЕХНОЛОГИЧЕСКАЯ КАРТА" if item["kind"]=="TTK"
         else "ТЕХНОЛОГИЧЕСКАЯ КАРТА")+" № "+str(item["id"])+" / ред. "+str(item["version"]),title),
      p(data["title"],title),
      p("Основание рецептуры: "+str(data.get("source_reference") or
                                "не указано — требуется уточнение"),small),
      p("Дата создания: "+item["created_at"][:10]+"; дата изменения: "+item["updated_at"][:10],small),
    ]
    for name,text in _sections(data):
        story.append(p(name,heading))
        if name.startswith("3."):
            rows=[[p("Наименование продукта",small),p("Брутто, г",small),p("Нетто, г",small)]]
            for ing in data["ingredients"]:
                rows.append([p(ing["name"],small),p(ing["gross_g"],small),p(ing["net_g"],small)])
            table=Table(rows,colWidths=[112*mm,30*mm,30*mm],repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#DCE8F5")),
                ("LINEBELOW",(0,0),(-1,-1),.35,colors.HexColor("#C4D1DC")),
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("TOPPADDING",(0,0),(-1,-1),4),
                ("BOTTOMPADDING",(0,0),(-1,-1),4),
            ]))
            story.extend([table,Spacer(1,5),
                p("Плановый выход на порцию: "+str(data.get("output_g",""))+
                  " г."+(" Требует подтверждения." if data.get("output_is_estimate") else ""),small)])
        else:
            story.append(p(text))
    story.extend([
       Spacer(1,13),
       p("ПОДГОТОВКА, СОГЛАСОВАНИЕ И УТВЕРЖДЕНИЕ",heading),
       p("Составил: "+(item.get("authored_by") or data.get("prepared_by") or "не указано")+
         "      Дата: ______________      Подпись: ______________",small),
       p("Проверил: "+(item.get("checked_by") or data.get("checked_by") or "не указано")+
         "      Дата: ______________      Подпись: ______________",small),
       p("Утверждающее лицо: "+(item.get("approved_by") or "не указано"),small),
       p("Дата подтверждения в программе: "+(item.get("approved_at") or "не утверждено"),small),
       p("Собственноручная подпись утверждающего: __________________________",small),
       p("ФИО и запись в программе не заменяют юридически значимую подпись.",small),
       p("Статус: "+("НЕУТВЕРЖДЁННЫЙ ПРОЕКТ" if draft else "ПОДТВЕРЖДЁН В СИСТЕМЕ")+
         ". ГОСТ 31987-2012, СанПиН 2.3/2.4.4282-26.",small),
    ])
    doc.build(story)
    return output


def as_docx(output,venue,item):
    from docx import Document
    from docx.shared import Cm,Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    data,draft,status=_common(venue,item)
    output=Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    document=Document()
    sec=document.sections[0]
    sec.top_margin=Cm(1.6);sec.bottom_margin=Cm(1.6)
    sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
    st=document.styles["Normal"]
    st.font.name="Arial";st.font.size=Pt(9)
    document.add_paragraph(str(venue.get("legal_name") or venue.get("venue_name") or ""))
    document.add_paragraph("Объект: "+str(venue.get("venue_name") or ""))
    document.add_paragraph("Адрес: "+str(venue.get("address") or "не указан"))
    p=document.add_paragraph(status)
    p.style="Subtitle"
    document.add_heading(("ТЕХНИКО-ТЕХНОЛОГИЧЕСКАЯ" if item["kind"]=="TTK"
                          else "ТЕХНОЛОГИЧЕСКАЯ")+" КАРТА № "+
                          str(item["id"])+" / ред. "+str(item["version"]),level=1)
    document.add_heading(data["title"],level=2)
    document.add_paragraph("Основание рецептуры: "+
       (data.get("source_reference") or "не указано — требуется подтверждение"))
    document.add_paragraph("Создана: "+item["created_at"][:10]+
                           "; изменена: "+item["updated_at"][:10])
    for name,text in _sections(data):
        document.add_heading(name,level=2)
        if name.startswith("3."):
            t=document.add_table(rows=1,cols=3)
            t.style="Light Shading Accent 1"
            for col,s in enumerate(("Наименование продукта","Брутто, г","Нетто, г")):
                t.rows[0].cells[col].text=s
            for ing in data["ingredients"]:
                row=t.add_row()
                for i,v in enumerate((ing["name"],str(ing["gross_g"]),str(ing["net_g"]))):
                    row.cells[i].text=v
            document.add_paragraph("Плановый выход на порцию: "+
              str(data.get("output_g"))+" г"+
              (" (ТРЕБУЕТ ПОДТВЕРЖДЕНИЯ)" if data.get("output_is_estimate") else ""))
        else:
            document.add_paragraph(text or "Не заполнено")
    document.add_heading("Подготовка и утверждение",level=2)
    document.add_paragraph("Составил: "+(item.get("authored_by") or data.get("prepared_by") or "не указано"))
    document.add_paragraph("Дата: _______________ Подпись: _____________________")
    document.add_paragraph("Проверил: "+(item.get("checked_by") or data.get("checked_by") or "не указано"))
    document.add_paragraph("Дата: _______________ Подпись: _____________________")
    document.add_paragraph("Утвердил: "+(item.get("approved_by") or "не утверждено"))
    document.add_paragraph("Дата записи об утверждении: "+(item.get("approved_at") or "—"))
    document.add_paragraph("Подпись утверждающего: _________________________")
    document.add_paragraph("ФИО и запись в программе не заменяют юридически значимую подпись.")
    document.add_paragraph("Статус: "+("НЕУТВЕРЖДЁННЫЙ ПРОЕКТ" if draft
                                          else "ПОДТВЕРЖДЁН В СИСТЕМЕ")+
                           ". ГОСТ 31987-2012 / СанПиН 2.3/2.4.4282-26.")
    document.save(str(output))
    return output
